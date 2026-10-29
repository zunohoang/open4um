from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
from datetime import datetime,timezone
from zoneinfo import ZoneInfo
import hashlib,json,re,argparse
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_TAB_ALIGNMENT,WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
parser=argparse.ArgumentParser(description='Rebuild the professor report from previously captured local evidence; no live test calls.')
parser.add_argument('--work-dir',type=Path,default=ROOT/'tmp/btl-tester')
parser.add_argument('--template',type=Path,required=True)
args=parser.parse_args()
OUT=args.work_dir.resolve()
QA=ROOT
REF=args.template.resolve()
reference_hash=hashlib.sha256(REF.read_bytes()).hexdigest()
assert reference_hash=='eee6ed0cdae50bf3eb21f37209cdbf7a4f5254e2f0cb9fb576daebd803698146'
UI_DIR=QA/'tests/selenium/reports/422a0a7242d44f1c852e96c9115e2031'
API_DIR=QA/'tests/postman/reports/643604ab509e4faba7cb0a064eed4b79'
UI=json.loads((UI_DIR/'summary.json').read_text(encoding='utf-8'))
API=json.loads((API_DIR/'summary.json').read_text(encoding='utf-8'))
AV=json.loads((OUT/'avatar/summary.json').read_text(encoding='utf-8'))
FIX=json.loads((OUT/'avatar/fixtures.json').read_text(encoding='utf-8'))
def check_integrity(directory):
    manifest=json.loads((directory/'integrity.json').read_text(encoding='utf-8'))
    for entry in manifest['files']:
        f=(directory/entry['path']).resolve()
        assert f.is_relative_to(directory.resolve()) and f.is_file()
        assert hashlib.sha256(f.read_bytes()).hexdigest()==entry['sha256']
check_integrity(UI_DIR);check_integrity(API_DIR)
assert len(UI['tests'])==30 and len(API['requests'])==15
assert sum(x['assertions'] for x in API['requests'])==48
assert AV['headerSynchronized'] and AV['cleanup']=='verified'

doc=Document(REF)
body=doc._element.body
cover_frame=deepcopy(list(body)[0])
cover_logo=deepcopy(list(body)[5])
for child in list(body):
    if child.tag!=qn('w:sectPr'):body.remove(child)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Times New Roman')
normal=doc.styles['Normal'];normal.font.size=Pt(13)
normal.paragraph_format.line_spacing=1.1
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.first_line_indent=Cm(0)
normal.paragraph_format.widow_control=True
for level,size in [(1,16),(2,14),(3,13)]:
    st=doc.styles[f'Heading {level}'];st.font.size=Pt(size);st.font.bold=True
    st.paragraph_format.space_before=Pt(12 if level>1 else 0)
    st.paragraph_format.space_after=Pt(6);st.paragraph_format.keep_with_next=True
    st.paragraph_format.keep_together=True
    pr=st.element.get_or_add_pPr()
    for n in list(pr):
        if n.tag==qn('w:numPr'):pr.remove(n)
doc.styles['Title'].font.size=Pt(22)
doc.styles['Title'].font.bold=True
doc.styles['Caption'].font.size=Pt(11)
doc.styles['Caption'].font.italic=True
doc.styles['Caption'].paragraph_format.space_after=Pt(8)
doc.core_properties.title='Báo cáo bài tập lớn kiểm thử phần mềm ABSlider'
doc.core_properties.author='Nguyễn Thái Duy'
doc.core_properties.last_modified_by='Nguyễn Thái Duy'
doc.core_properties.created=datetime.now(timezone.utc)
doc.core_properties.modified=datetime.now(timezone.utc)
doc.core_properties.revision=1
doc.core_properties.subject='Phân tích thiết kế và kết quả kiểm thử ba chức năng ABSlider'
doc.core_properties.keywords='TC-9; TC-73; TC-74; Selenium; Postman; AI outline'
doc.core_properties.comments='Báo cáo dựa trên bằng chứng thực thi; không thay thế phê duyệt phát hành.'
settings=doc.settings.element
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)

def p(text='',bold=False,style=None):
    para=doc.add_paragraph(style=style)
    para.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold:para.paragraph_format.keep_with_next=True
    run=para.add_run(text);run.bold=bold
    return para

def h(text,level=2):
    para=doc.add_paragraph(text,f'Heading {level}')
    para.alignment=WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.first_line_indent=Cm(0)
    return para

def chapter(text):
    doc.add_page_break();return h(text,1)

table_no=0
def table(headers,rows,widths):
    global table_no
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'7F7F7F');borders.append(el)
    props.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side,val in [('top',80),('bottom',80),('left',95),('right',95)]:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');margins.append(e)
    props.append(margins)
    for i,text in enumerate(headers):t.rows[0].cells[i].text=text
    header=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');header.append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for i,text in enumerate(row):cells[i].text=str(text)
    for ri,row in enumerate(t.rows):
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'1F497D');cell._tc.get_or_add_tcPr().append(sh)
            for para in cell.paragraphs:
                para.paragraph_format.first_line_indent=Cm(0);para.paragraph_format.space_after=Pt(3)
                para.paragraph_format.line_spacing=1.0;para.paragraph_format.keep_with_next=False
                para.alignment=WD_ALIGN_PARAGRAPH.LEFT
                for r in para.runs:
                    r.font.name='Times New Roman';r.font.size=Pt(11.5);r.bold=ri==0
                    r.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def caption_table(title,headers,rows,widths):
    global table_no
    table_no+=1
    q=doc.add_paragraph(f'Bảng {table_no}. {title}','Caption');q.paragraph_format.keep_with_next=True
    table(headers,rows,widths)

fig_no=0
def figure(path,title,claim,maxheight=15,crop=None):
    global fig_no
    from PIL import Image
    im=Image.open(path);w,hp=im.size
    visible_w,visible_h=(crop[2],crop[3]) if crop else (w,hp)
    width=min(16.0,maxheight*visible_w/visible_h)
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next=True
    picture=para.add_run().add_picture(str(path),width=Cm(width))
    if crop:
        x,y,cw,ch=crop
        picture.height=Cm(width*ch/cw)
        src=OxmlElement('a:srcRect')
        for key,value in [('l',x/w),('t',y/hp),('r',(w-x-cw)/w),('b',(hp-y-ch)/hp)]:src.set(key,str(round(value*100000)))
        fill=picture._inline.graphic.graphicData.pic.blipFill
        fill.insert(1,src)
    fig_no+=1
    q=doc.add_paragraph(f'Hình {fig_no}. {title}','Caption');q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.keep_with_next=True
    p(claim)

def code(text):
    q=doc.add_paragraph();q.paragraph_format.line_spacing=1.0
    q.paragraph_format.space_after=Pt(8)
    r=q.add_run(text);r.font.name='Consolas';r.font.size=Pt(10.5)
    return q

def localtime(value):return datetime.fromisoformat(value).astimezone(ZoneInfo('Asia/Ho_Chi_Minh')).strftime('%d/%m/%Y %H:%M:%S')

# Retain the actual frame and logo, not just a loosely similar cover.
body.insert(0,cover_frame)
# Cover derived from the reference's university/topic/author/date hierarchy.
for text,bold,size,space in [
 ('TRƯỜNG ĐẠI HỌC GIAO THÔNG VẬN TẢI',True,13,4),('KHOA CÔNG NGHỆ THÔNG TIN',True,13,2),('----o0o----',False,13,18),
 ('BÁO CÁO BÀI TẬP LỚN',True,18,8),('KIỂM THỬ PHẦN MỀM',True,18,18),
 ('ĐỀ TÀI',True,14,10),('XÂY DỰNG WEBSITE TẠO SLIDE',True,15,6),('BÀI GIẢNG BẰNG AI',True,15,24),
 ('Dự án ABSlider — Nhóm 03',False,13,24),
 ('Sinh viên thực hiện: Nguyễn Thái Duy',True,13,8),('Mã sinh viên: 231210730',False,13,24),
 ('Ba chức năng: ảnh đại diện, cấu hình Credit, sinh dàn ý AI',False,12,28),('Hà Nội, tháng 10 năm 2026',False,13,0)]:
    q=doc.add_paragraph(style='Title' if text in ['BÁO CÁO BÀI TẬP LỚN','KIỂM THỬ PHẦN MỀM'] else 'Normal')
    q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.space_after=Pt(space)
    r=q.add_run(text);r.bold=bold;r.font.size=Pt(size)
    if text=='----o0o----':
        body.insert(len(body)-1,cover_logo)

doc.add_page_break();h('MỤC LỤC',1)
toc_marker=doc.add_paragraph()
p('Báo cáo trình bày kế hoạch, căn cứ thiết kế và bằng chứng kiểm thử ba chức năng của ABSlider. Trọng tâm là giải thích vì sao chọn dữ liệu và cách xác định kết quả; Selenium và Postman là công cụ hỗ trợ thực hiện các điều kiện kiểm thử đã được phân tích.')
p('Phạm vi gồm một ca ảnh đại diện, hai ca cấu hình Credit và ba kịch bản API sinh dàn ý. Các con số về ca, biến thể dữ liệu, yêu cầu HTTP và phép khẳng định được tách riêng để tránh nhầm độ lớn bộ kiểm thử với mức độ bao phủ chức năng.')

chapter('1. LẬP KẾ HOẠCH KIỂM THỬ')
h('1.1. Mục tiêu và đối tượng')
p('ABSlider hỗ trợ người dùng tạo bài giảng bằng AI, quản lý hồ sơ và sử dụng Credit cho các thao tác AI. Báo cáo lựa chọn ba chức năng có quy tắc đầu vào rõ, có trạng thái lưu trữ cần đối chiếu và có thể thu thập bằng chứng thực tế. Mục tiêu là xác định dữ liệu nào được chấp nhận hoặc bị từ chối, thay đổi nào được lưu và kết quả quan sát có đáp ứng đúng yêu cầu đã chọn hay không.')
p('Nguyễn Thái Duy chịu trách nhiệm phân tích, thiết kế ca, tổ chức thực hiện và tổng hợp báo cáo. Đây là phạm vi kiểm thử của bài tập, không phải kết luận chất lượng cho toàn bộ 32 Use Case. Các kết quả tự động được tổng hợp từ hai đợt chạy đã lưu; TC-9 được kiểm tra bổ sung ngày 29/10/2026. Báo cáo không giả định rằng bản kế hoạch này đã được phê duyệt trước các đợt chạy đó.')
h('1.2. Phạm vi và lựa chọn phương pháp')
caption_table('Ba chức năng và phương pháp kiểm thử',['Chức năng','Ca và nội dung chính','Phương pháp'],[
 ('Ảnh đại diện','TC-9 / TC-UC004-02: định dạng, dung lượng, preview, lưu và đồng bộ Header.','Thiết kế thủ công theo bước; đợt xác minh dùng trình duyệt có trợ giúp WebDriver.'),
 ('Cấu hình Credit','TC-73 / UC032-01: lưu giá hợp lệ. TC-74 / UC032-02: chặn giá không hợp lệ, sửa lại và kiểm tra không lưu một phần.','Tự động bằng Python, pytest và Selenium; 30 biến thể thực thi.'),
 ('Sinh dàn ý AI','AI-OUTLINE-01: prompt hợp lệ. 02: prompt rỗng. 03: thiếu xác thực.','Tự động bằng Postman; kiểm tra phản hồi và đọc lại dữ liệu độc lập.')],[3.0,8.0,5.3])
p('Ảnh đại diện cần đối chiếu trực quan thông báo, khung xem trước và Header nên được thiết kế như một ca kiểm tra giao diện theo bước. Cấu hình Credit có nhiều đầu vào lặp lại, phù hợp với tự động hóa để giữ cùng cách nhập và kiểm tra sau mỗi lần làm mới. Với sinh dàn ý, API cho phép quan sát trực tiếp mã HTTP, cấu trúc dữ liệu, số Credit bị trừ và dữ liệu đã lưu, tránh phụ thuộc giao diện khi đánh giá hợp đồng backend.')
p('Không đưa vào phạm vi: UC032-03 về cập nhật giá giữa nhiều phiên; sinh toàn bộ slide; AI Chat/chỉnh sửa; email/OTP; toàn bộ CRUD backend; kiểm thử tải, hiệu năng và đánh giá chất lượng sư phạm của nội dung AI. Collection Postman toàn backend được giữ để tham khảo nhưng không chạy trong đợt API được báo cáo.')
h('1.3. Phân công trách nhiệm')
caption_table('Trách nhiệm và đầu ra',['Vai trò','Người phụ trách','Công việc và đầu ra'],[
 ('Tester và người lập báo cáo','Nguyễn Thái Duy — 231210730','Chọn yêu cầu, phân tích rủi ro, thiết kế dữ liệu/tiêu chí đạt, đối chiếu bằng chứng và tổng hợp kết quả của cả ba chức năng.'),
 ('Hỗ trợ thao tác và tự động hóa','Trợ lý AI cùng bộ công cụ chạy thử','Thực hiện thao tác trình duyệt có trợ giúp, chạy bộ tự động đã được cho phép và lưu bằng chứng. Không ghi nhận thành thao tác thủ công trực tiếp của Duy.'),
 ('Hỗ trợ phát triển','Nhóm phát triển Group 03','Hỗ trợ môi trường, phân tích và sửa lỗi khi được chuyển giao. Báo cáo không tự gán thành viên vào một lượt kiểm thử chưa có bằng chứng.')],[3.2,4.6,8.5])
h('1.4. Môi trường và dữ liệu')
caption_table('Môi trường thực thi được ghi nhận',['Thành phần','Giá trị và cách sử dụng'],[
 ('Môi trường','Hosted development; frontend https://dev-slides.sbltcup.dev và API https://dev-slides-api.sbltcup.dev/api/v1.'),
 ('Bản backend cho hai lượt tự động','develop, SHA fdaa8c44933b3737c144a023b995de8f54848a05; lấy từ metadata lượt chạy, không suy ra từ tên nhánh.'),
 ('Trình duyệt và công cụ','Chrome '+str(UI['tests'][0].get('browser',{}).get('version','theo metadata'))+' trên Windows; Python 3.12, Selenium 4.38, pytest 8.4.2; Postman CLI 1.82.0.'),
 ('Tài khoản','Admin QA dùng thiết lập/khôi phục cấu hình và tạo fixture. API outline và TC-9 dùng người dùng thử riêng; mật khẩu và token không đưa vào báo cáo.'),
 ('Credit','Mốc kiểm thử giao diện 80/10/5/5/100. API đọc giá outline đang áp dụng; không ghi đè cấu hình. Các lượt có thể thay đổi giá phải thực hiện tuần tự.'),
 ('Phụ thuộc','Dịch vụ backend, dữ liệu và xác thực phải sẵn sàng. Lượt AI cần Gemini hoạt động; lỗi phụ thuộc không được chuyển thành kết quả đạt.')],[4.0,12.3])
h('1.5. Trình tự và điều kiện hoàn thành')
p('Trình tự công việc: xác nhận yêu cầu và trạng thái ca; chuẩn bị tài khoản/tệp/mốc Credit; thiết kế các điều kiện và kết quả mong đợi; thực hiện theo từng chức năng; đối chiếu quan sát với tiêu chí; thu hồi dữ liệu thử; tổng hợp lỗi và giới hạn. Không chạy đồng thời các bộ có thể tác động tới cấu hình Credit chung.')
caption_table('Điều kiện bắt đầu và kết thúc',['Mốc','Điều kiện'],[
 ('Bắt đầu','Xác nhận đúng môi trường develop, tài khoản đúng quyền, dữ liệu thử đã đo được, tiêu chí đạt đã tách theo bước và có nơi lưu bằng chứng.'),
 ('PASSED','Tất cả điều kiện bắt buộc của ca có bằng chứng đạt; dữ liệu được lưu hoặc bị chặn đúng; hậu kiểm/thu hồi thành công.'),
 ('FAILED','Ít nhất một quan sát trái yêu cầu hoặc hậu kiểm thất bại. Không lấy một bước thành công để thay cho toàn ca.'),
 ('BLOCKED / NOT RUN','Không thể thực hiện vì thiếu điều kiện, hoặc chưa có lần thực thi. CONFIRMED chỉ xác nhận nội dung ca đã được xem xét, không phải kết quả chạy.'),
 ('Hoàn thành báo cáo','Đủ ma trận truy vết, kết quả thực tế, bằng chứng chọn lọc và danh sách sai lệch; các phần chưa kiểm tra phải được chỉ rõ.')],[3.5,12.8])
h('1.6. Rủi ro và biện pháp kiểm soát')
p('Rủi ro ưu tiên là thay đổi giá Credit chung, ảnh chụp sau khi giao diện đã bị reset, AI trả nội dung không cố định và nhầm “đã có script” với “đã chạy đạt”. Biện pháp tương ứng là chạy tuần tự và lưu/khôi phục mốc cũ; chụp tại thông báo hoặc trạng thái vừa phát sinh; dùng tiêu chí cấu trúc và đối chiếu Credit thay vì so nguyên văn nội dung AI; gắn mọi kết quả với mã lượt chạy và nguồn bằng chứng. Khóa chạy trong một checkout không ngăn được Admin trên máy khác, nên vẫn cần tránh cửa sổ thay đổi đồng thời.')

chapter('2. PHÂN TÍCH VÀ THIẾT KẾ KIỂM THỬ')
h('2.1. Cách chuyển yêu cầu thành điều kiện kiểm thử')
p('Một chức năng được phân tích theo đầu vào, quy tắc xử lý, trạng thái đầu ra và rủi ro nếu xử lý sai. Từ đó xác định điều kiện cần kiểm tra, chọn đại diện cho mỗi lớp dữ liệu và xác định nguồn đối chiếu độc lập. Các bảng thiết kế dưới đây thể hiện việc chọn dữ liệu có chủ đích; số lượng công cụ hoặc số lần gửi request không thay thế cho lập luận này.')
p('Nguồn yêu cầu trực tiếp là TC-9 do người dùng cung cấp, định nghĩa Sprint 4 của TC-73/74 và ba kịch bản outline đã thống nhất. Mã nguồn hỗ trợ giải thích ràng buộc kỹ thuật nhưng không được dùng để sửa ngầm kết quả mong đợi cho khớp với phần mềm đang chạy. Khi khác nhau, báo cáo ghi nhận sai lệch.')
h('2.2. Chức năng tải ảnh đại diện')
h('2.2.1. Quy tắc nghiệp vụ và rủi ro',3)
p('Người dùng tại /profile được chọn ảnh đại diện mới. Hệ thống phải từ chối tệp sai định dạng hoặc vượt giới hạn, không đưa tệp đó vào preview. Ảnh hợp lệ phải xem trước ngay, được lưu khi bấm “Lưu thay đổi” và được đồng bộ lên Header. TC-9 còn yêu cầu khung preview tròn và thông báo từ chối đúng nội dung đã quy định. Vì vậy, “upload thành công” chỉ là một phần của tiêu chí đạt.')
caption_table('Phân tích yêu cầu ảnh đại diện',['Mã','Điều kiện và rủi ro','Cách xác định đúng/sai'],[
 ('AV-01','Từ chối PDF; tránh đưa tệp không phải ảnh vào trạng thái hồ sơ.','Thông báo đúng yêu cầu; không xuất hiện preview mới từ PDF; không phát sinh lưu hồ sơ chỉ do chọn tệp.'),
 ('AV-02','Từ chối ảnh lớn hơn 2 MB; tránh dữ liệu quá giới hạn.','Ảnh JPEG có định dạng hợp lệ nhưng kích thước vượt mốc; có lỗi, không có preview mới.'),
 ('AV-03','Xem trước PNG hợp lệ; tránh chọn tệp được nhưng người dùng không nhận ra ảnh nào sẽ lưu.','Preview là ảnh vừa chọn và nằm trong khung tròn như TC; chưa đồng nhất preview với việc đã lưu.'),
 ('AV-04','Lưu và đồng bộ; tránh trạng thái chỉ thay đổi trên form.','Thông báo lưu thành công, Header dùng cùng ảnh, GET /auth/me trả ảnh tương ứng và tải lại vẫn còn.')],[1.5,7.1,7.7])
h('2.2.2. Phân lớp và phân tích giá trị biên',3)
caption_table('Lớp dữ liệu và giá trị biên của tệp',['Khía cạnh','Lớp/giá trị','Lý do và phạm vi'],[
 ('Định dạng','Ảnh được hỗ trợ / tệp PDF','PDF nhỏ được chọn để cô lập lỗi định dạng, không trộn với lỗi kích thước.'),
 ('Dung lượng hợp lệ','0 < kích thước < 2.097.152 byte','PNG nhỏ là đại diện cho đường thành công. Không suy ra mọi ảnh nhỏ đều hợp lệ nếu tệp hỏng.'),
 ('Biên đúng bằng giới hạn','2.097.152 byte','Mã hiện tại dùng điều kiện file.size > giới hạn nên đúng bằng mốc được dự kiến chấp nhận; đây là phân tích biên, chưa thực thi trong ba tệp đã chọn.'),
 ('Dung lượng không hợp lệ','Kích thước > 2.097.152 byte','Dùng JPEG thực có thể đọc được để loại trừ lỗi “tệp giả ảnh”; xác minh lỗi dung lượng.'),
 ('Ngoài phạm vi lần chạy','Tệp 0 byte, ảnh hỏng, đổi đuôi, ảnh nén bất thường','Không có bằng chứng thực thi; không gộp thành kiểm thử bảo mật hay bao phủ đầy đủ mọi định dạng.')],[3.0,5.5,7.8])
p('Báo cáo dùng cách tính của phần mềm là 2 × 1024 × 1024 byte. Không dùng kích thước làm tròn trên giao diện quản lý tệp để kết luận ở biên. Khi từ chối tệp mới, “không preview” nghĩa là không thay preview bằng tệp lỗi; nếu người dùng vốn có avatar thì ảnh cũ có thể vẫn hiển thị.')
caption_table('Dữ liệu thực tế cho TC-9',['Tệp','Kích thước byte','Mục đích'],[(x['name'],format(x['bytes'],','),'Sai định dạng, nhỏ hơn mốc' if x['name'].endswith('.pdf') else 'Ảnh JPEG hợp lệ nhưng vượt mốc' if x['name'].endswith('.jpg') else 'Ảnh PNG hợp lệ để xem trước và lưu') for x in sorted(FIX,key=lambda x:x['name'])],[5.0,3.5,7.8])
h('2.2.3. Thiết kế TC-9 và tiêu chí đánh giá',3)
p('Tiền điều kiện: đã đăng nhập bằng người dùng thử, mở /profile, ghi nhận avatar ban đầu; dùng tệp đúng tên và kích thước đã chuẩn bị. Không thao tác sửa họ tên trong ca này để cô lập thay đổi avatar. Nếu chọn tệp lỗi, không bấm lưu như một ca sửa hồ sơ khác.')
caption_table('Thiết kế theo bước của TC-9',['Bước','Thao tác và dữ liệu','Kết quả mong đợi / bằng chứng'],[
 ('1','Chọn document.pdf.','Hiển thị đúng thông báo quy định; không preview PDF. Ảnh chụp phải thấy lỗi và khu vực preview.'),
 ('2','Chọn avatar_heavy.jpg.','Báo vượt dung lượng theo thông báo quy định; không preview ảnh lớn. Ghi kích thước tệp đã chọn.'),
 ('3a','Chọn avatar_valid.png.','Hiện ngay preview ảnh trong khung tròn; không coi ảnh preview là ảnh đã lưu.'),
 ('3b','Bấm Lưu thay đổi.','Có phản hồi thành công, Header đồng bộ cùng ảnh. Lưu ảnh chụp trước khi tải lại.'),
 ('Hậu kiểm','Tải lại /profile và mở lại workspace; GET /auth/me.','Ảnh không mất; Header và dữ liệu hồ sơ cùng tham chiếu ảnh mới. Đọc API chỉ hỗ trợ đối chiếu, không thay thao tác lưu trên UI.')],[1.6,5.4,9.3])
p('Thông báo mong đợi giữ nguyên: “Ảnh vượt quá dung lượng cho phép (tối đa 2MB) hoặc sai định dạng tệp tin”. Báo cáo đối chiếu nội dung thông báo, đồng thời phân biệt khác biệt dấu chấm kết câu với việc thay cả nội dung hoặc thiếu thông tin. Sai lệch giao diện/thông điệp không bị bỏ qua chỉ vì thao tác từ chối hoặc lưu đã thành công.')
h('2.3. Chức năng cấu hình Credit')
h('2.3.1. Mô hình dữ liệu và tính nhất quán',3)
p('Admin quản lý bốn mức tiêu hao AI và Credit tặng khi đăng ký. Đổi giá sai trường có thể khiến thao tác AI bị tính sai; lưu một phần khi form có lỗi có thể làm hệ thống rơi vào trạng thái khó nhận biết. Vì thế, kiểm tra chỉ nhìn thông báo “thành công” là chưa đủ: phải đối chiếu toàn bộ cấu hình sau F5 và qua lần đọc backend độc lập.')
caption_table('Ánh xạ trường Credit và mốc kiểm thử',['Nhãn nghiệp vụ','Tên API','Mốc'],[
 ('Giá mỗi lượt sinh bài giảng','pricePerLecture','80'),('Giá mỗi slide tạo mới trong AI Chat','pricePerSlide','10'),('Giá mỗi lượt AI chỉnh sửa slide','pricePerAiEdit','5'),('Giá mỗi lượt AI sinh dàn ý','pricePerOutline','5'),('Credit tặng khi đăng ký mới','signupBonus','100')],[7.5,6.3,2.5])
p('Mốc 80/10/5/5/100 là dữ liệu điều khiển của hai ca, không phải tuyên bố mọi bản triển khai buộc có giá đó. Trước lượt chạy, bộ kiểm thử ghi cấu hình thực tế; trước mỗi biến thể thiết lập mốc; sau cùng khôi phục giá trị ban đầu của môi trường. Trong phạm vi invalid của TC-74, bốn giá AI được thay đổi còn signupBonus giữ 100.')
h('2.3.2. TC-73 lưu cấu hình hợp lệ',3)
caption_table('Chuỗi trạng thái hợp lệ trong TC-73',['Lần lưu','Thay đổi','Toàn bộ cấu hình phải thấy sau F5'],[
 ('Ban đầu','Đọc mốc đã lưu','80 / 10 / 5 / 5 / 100'),('1','AI Chat: 10 → 7','80 / 7 / 5 / 5 / 100'),('2','AI edit: 5 → 3','80 / 7 / 3 / 5 / 100'),('3','Sinh bài giảng: 80 → 81','81 / 7 / 3 / 5 / 100'),('4','Sinh dàn ý: 5 → 6','81 / 7 / 3 / 6 / 100'),('5','Biên AI Chat: 7 → 1','81 / 1 / 3 / 6 / 100'),('Khôi phục','Đặt lại mốc','80 / 10 / 5 / 5 / 100')],[2.3,6.0,8.0])
p('Mỗi lần lưu phải thỏa bốn điều kiện: thông báo rõ ràng khi lưu; đúng trường nhận giá mới; các trường khác giữ nguyên; dữ liệu vẫn còn sau refresh. Giá 1 đại diện cho biên dưới hợp lệ của số nguyên dương. Không có cơ sở để tuyên bố kiểm tra hết miền số hoặc giá trị cực đại khi chỉ thực thi những giá này.')
h('2.3.3. TC-74 chặn dữ liệu không hợp lệ',3)
caption_table('Phân lớp đầu vào của TC-74',['Lớp','Đại diện','Rủi ro cần phát hiện'],[
 ('Không dương tại biên','0','Vô hiệu hóa chi phí hoặc vi phạm quy tắc số nguyên dương.'),('Số âm','-1','Tính toán sai hoặc làm tăng số dư khi dùng như một mức trừ.'),('Không nguyên','1.5','Nhận giá lẻ dù yêu cầu là số nguyên.'),('Không phải số','abc','Chuyển đổi sai thành một số hoặc giá rỗng.'),('Trộn số và chữ','5abc','Cắt bỏ phần chữ rồi lưu nhầm 5.'),('Trống','Xóa toàn bộ giá trị','Âm thầm coi rỗng là 0 hoặc bỏ qua kiểm tra bắt buộc.'),('Khoảng trắng','Ba dấu cách','Không trim/không kiểm tra dẫn tới dữ liệu không hợp lệ.')],[3.8,3.8,8.7])
p('Bốn trường AI × bảy lớp dữ liệu tạo thành 28 biến thể độc lập. Mỗi biến thể chỉ thay một trường, thử lưu, kiểm tra lỗi hoặc hành vi chặn nhập, rồi tải lại và đối chiếu toàn bộ năm trường với mốc. Trường HTML kiểu number có thể không cho ký tự đi vào ô; bằng chứng phải ghi cả chuỗi định nhập và giá trị thực tế trình duyệt nhận, không tuyên bố đã gửi nguyên chuỗi “abc” tới API.')
p('Biến thể kết hợp thứ 29 của TC-74 đặt giá tạo slide bằng 6 và giá chỉnh sửa bằng 0. Trước khi sửa, không trường nào được lưu một phần. Sau khi đổi giá chỉnh sửa thành 6, lưu và F5 phải cho 80/6/6/5/100. Cuối cùng khôi phục mốc. Trường hợp này kiểm tra đường phục hồi từ dữ liệu sai về dữ liệu đúng, không chỉ riêng thông báo từ chối.')
caption_table('Điều kiện bắt buộc và nguồn đối chiếu Credit',['Điều kiện','Quan sát cần có','Nguồn bằng chứng'],[
 ('Dữ liệu sai bị chặn','Không thông báo thành công; trường/bảng lỗi giải thích số nguyên dương.','Ảnh tại thời điểm validation và giá trị ô nhập thực tế.'),('Không gửi/lưu sai','Không có PATCH Credit trong các lần UI chặn; sau F5 dữ liệu không đổi.','Giao dịch trình duyệt được lọc theo Credit API; GET backend độc lập.'),('Sửa đúng có thể lưu','Form hết lỗi, lưu được, dữ liệu còn sau refresh.','Ảnh save-feedback, persisted checkpoint và API readback.'),('Không ảnh hưởng trường khác','Tất cả năm trường khớp trạng thái mong đợi.','Đối chiếu tập giá trị thay vì chỉ kiểm tra ô vừa sửa.')],[3.5,6.1,6.7])
h('2.4. Chức năng sinh dàn ý AI qua API')
h('2.4.1. Hợp đồng và rủi ro',3)
p('Endpoint được chọn là POST /lectures/generate-outline. Backend nhận prompt và trả dữ liệu dàn ý cùng thông tin bài giảng/Credit. Phiên bản đang kiểm tra không có tham số slide-count riêng. Do AI không bảo đảm câu chữ cố định, tiêu chí đạt dựa vào cấu trúc có thể sử dụng: tiêu đề không rỗng, danh sách sections không rỗng, mỗi section có heading và mảng bullets với phần tử dạng chuỗi.')
p('Ba rủi ro ưu tiên là trả kết quả không dùng được dù HTTP 200; chấp nhận đầu vào thiếu nội dung; và cho người chưa đăng nhập sử dụng chức năng AI. Với đường thành công, kiểm tra thêm số dư và outline đã lưu để phát hiện trường hợp phản hồi đúng nhưng dữ liệu backend chưa được cập nhật. Các tiêu chí này không chứng minh nội dung đúng về học thuật hoặc toàn bộ mọi loại prompt.')
caption_table('Thiết kế ba kịch bản API',['Kịch bản','Đầu vào/tiền điều kiện','Kết quả mong đợi'],[
 ('AI-OUTLINE-01','User thử có số dư đủ; prompt yêu cầu giải thích phép cộng số tự nhiên, ví dụ 2+3=5 và 4+1=5.','HTTP 200; success=true; cấu trúc outline hợp lệ; creditSpent bằng giá outline; số dư giảm đúng; GET hồ sơ và bài giảng khớp phản hồi.'),
 ('AI-OUTLINE-02','Cùng user đã đăng nhập; body có prompt là chuỗi rỗng.','HTTP 422; success=false; có thông báo; không mất Credit. Kiểm tra số dư sau hai request từ chối, trước request thành công.'),
 ('AI-OUTLINE-03','Prompt có nội dung hợp lệ; không có Bearer token.','HTTP 401; success=false; có thông báo xác thực. Không coi HTTP 4xx mong đợi là lỗi của ca.')],[3.0,6.0,7.3])
h('2.4.2. Bố trí dữ liệu và đối chiếu Credit',3)
p('Một tài khoản thường dùng riêng cho lượt chạy được Admin tạo trước. Bộ kiểm thử đọc giá outline hiện hành, cấp cho fixture số Credit bằng giá đó cộng 10, rồi đăng nhập fixture. Với giá thực tế 5, số dư đầu là 15 và sau một lần sinh thành công phải là 10. Dữ liệu giá là đầu vào đọc từ môi trường, không được hard-code để ép phản hồi đạt.')
code('balance_after = balance_before - creditSpent\ncreditSpent = configured_price_per_outline\nGET /auth/me.creditBalance = balance_after\nGET /lectures/{id}.outline = generated_outline')
p('Chỉ một request có xác thực và prompt hợp lệ được gửi tới endpoint sinh dàn ý; không tự động thử lại khi lỗi. Hai request âm được thiết kế để dừng ở xác thực hoặc validation. Bộ kiểm thử không khẳng định quan sát được số lần gọi nội bộ từ backend tới Gemini; số lượng trên là request do bộ Postman chủ động phát sinh. Sau chạy, fixture được xóa bằng Admin API, đồng thời xóa dữ liệu sở hữu theo cơ chế backend.')
h('2.4.3. Phân biệt kịch bản và request hỗ trợ',3)
caption_table('Vai trò 15 request trong collection nhỏ',['Nhóm','Số HTTP request','Nội dung'],[
 ('Thiết lập','5','Kiểm tra develop, đăng nhập Admin, đọc giá outline, tạo một fixture, đăng nhập fixture.'),('Kịch bản và đọc lại','6','Thiếu xác thực; prompt rỗng; số dư sau từ chối; sinh thật; số dư sau sinh; bài giảng đã lưu.'),('Thu hồi','4','Đăng xuất fixture, xóa user cùng dữ liệu sở hữu, xác nhận user vắng mặt, đăng xuất Admin.')],[4.0,3.3,9.0])
p('Collection có 48 assertion trong lượt được báo cáo, nhưng đó không phải 48 ca độc lập. Một request có thể kiểm tra đồng thời HTTP status, envelope và nhiều thuộc tính. Request thiết lập/thu hồi giúp ca có dữ liệu sạch và bằng chứng đáng tin; chúng không được quảng bá thành chức năng nghiệp vụ bổ sung.')
h('2.5. Ma trận truy vết tổng hợp')
caption_table('Từ yêu cầu đến bằng chứng',['Yêu cầu','Ca / biến thể','Tiêu chí chính','Bằng chứng'],[
 ('AV-01, AV-02','TC-9 bước 1–2','Đúng lỗi; không preview tệp sai.','Ảnh PDF/JPEG lỗi, kích thước và hash fixture.'),('AV-03, AV-04','TC-9 bước 3 và hậu kiểm','Khung tròn; lưu; Header đồng bộ; còn sau F5.','Ảnh preview/save/Header; hash avatar từ GET /auth/me.'),('CR-01','TC-73: 1 chuỗi thực thi','Lưu giá hợp lệ, đúng trường, giữ sau F5.','Timeline của ca hợp lệ trong run Selenium.'),('CR-02','TC-74: 28 biến thể đơn','Số không hợp lệ không được lưu.','Ảnh validation + dữ liệu ô nhập + GET.'),('CR-03','TC-74: 1 biến thể kết hợp','Không lưu một phần; sửa đúng thì lưu.','Timeline correction và trạng thái 80/6/6/5/100.'),('AI-01','AI-OUTLINE-01','Cấu trúc, Credit và persistence.','API-909/910/911; phản hồi và assertion.'),('AI-02','AI-OUTLINE-02','422, thông báo, số dư không đổi.','API-907/908.'),('AI-03','AI-OUTLINE-03','401 khi không token.','API-906.')],[2.7,4.0,5.0,4.6])

chapter('3. THỰC HIỆN KIỂM THỬ')
h('3.1. Nhật ký và nguyên tắc sử dụng bằng chứng')
caption_table('Các lượt thực thi được đưa vào báo cáo',['Chức năng','Thời điểm GMT+7','Nhận diện lượt / kết quả'],[
 ('Credit — Selenium',localtime(UI['started']),'422a0a7242d44f1c852e96c9115e2031; 30/30 biến thể đạt; thu hồi xác minh.'),('Outline — Postman',localtime(API['started']),'643604ab509e4faba7cb0a064eed4b79; 15/15 request và 48/48 assertion đạt.'),('Avatar — trình duyệt hỗ trợ',localtime(AV['started']),AV['runId']+'; phát hiện hai sai lệch TC-9; fixture đã thu hồi.')],[4.0,4.0,8.3])
p('Hai bộ tự động dùng lại bằng chứng đã lưu đúng ngày thực thi; không chạy lại AI chỉ để có hình minh họa mới. TC-9 được thực hiện bổ sung qua một phiên Chrome được trợ lý điều khiển bằng WebDriver theo các bước đã thiết kế; đây không phải ghi nhận Duy trực tiếp thao tác thủ công. Trong báo cáo, “thủ công” mô tả phương pháp thiết kế/đối chiếu giao diện, còn cơ chế hỗ trợ thao tác của lượt này được nêu rõ.')
p('Ảnh trong báo cáo lấy từ checkpoint thực tế, không vẽ thêm thông báo lên ứng dụng. Các bảng actual result và đoạn JSON được trích từ receipt/summary đã lưu. Ảnh chọn lọc phục vụ đọc báo cáo; không thay toàn bộ timeline, JUnit và log giao dịch gốc. File manifest kiểm tra nội dung tệp so với bản đã lưu, không cung cấp chữ ký số hoặc chứng thực thời gian độc lập.')
h('3.2. Thực hiện TC-9')
caption_table('Quan sát thực tế của TC-9',['Bước','Kết quả quan sát','Đánh giá'],[
 ('PDF','Hiện: “Ảnh vượt quá dung lượng cho phép hoặc sai định dạng tệp tin (chỉ chấp nhận JPG, PNG, WEBP, GIF).”; không có preview mới.','Từ chối đúng; nội dung thông báo khác yêu cầu.'),('JPEG quá lớn','Hiện thông báo có “tối đa 2MB”; không có preview mới.','Đạt điều kiện từ chối kích thước; dấu chấm kết câu không đổi ý nghĩa.'),('PNG hợp lệ','Có preview đúng ảnh; khung vuông, border-radius của khung bằng 0px.','Preview tức thì đạt; hình dạng không đúng khung tròn.'),('Lưu và đọc lại','“Cập nhật hồ sơ thành công”; Header dùng ảnh mới; tải lại và GET /auth/me cùng ảnh fixture.','Đạt lưu/persistence/đồng bộ.'),('Thu hồi','Admin đối chiếu và xóa đúng user trong namespace của lượt; xác nhận user không còn.','Đã xác minh.')],[2.9,9.0,4.4])
figure(OUT/'avatar/01-pdf-rejected.png','TC-9 từ chối PDF nhưng thông báo khác yêu cầu','Bằng chứng AV-E01, vùng trích từ PNG gốc: banner lỗi và khu vực preview cùng xuất hiện; không có ảnh PDF mới. Sai lệch được giữ là lỗi hợp đồng thông báo, dù chức năng chặn tệp đã hoạt động.',10.0,(385,480,575,310))
figure(OUT/'avatar/02-heavy-rejected.png','TC-9 từ chối JPEG vượt 2 MB','Bằng chứng AV-E02, vùng trích từ PNG gốc: JPEG 5.685.841 byte bị chặn. Tệp là ảnh có định dạng hợp lệ; kích thước được ghi trong manifest thay vì suy ra từ ảnh màn hình.',10.0,(385,480,575,300))
figure(OUT/'avatar/04-save-success.png','TC-9 lưu PNG thành công và đồng bộ ảnh','Bằng chứng AV-E03: thông báo thành công hiển thị cùng avatar mới; góc trên bên phải dùng cùng ảnh. Khung preview vẫn vuông, tạo sai lệch thứ hai so với yêu cầu tròn.',14.3)
p('Đối chiếu độc lập cho thấy giá trị avatar do GET /auth/me trả về khớp data URL của PNG fixture; hash giá trị được ghi ở phụ lục. Ảnh sau refresh và ảnh workspace cũng được lưu trong hồ sơ bằng chứng. Việc Header hiển thị cùng ảnh được đối chiếu với cùng hash nguồn ảnh, không chỉ dựa vào màu giống nhau.')
h('3.3. Thực hiện hai ca Credit bằng Selenium')
p('TC-73 được chạy thành một chuỗi trạng thái hợp lệ; TC-74 gồm 28 biến thể dữ liệu độc lập và một biến thể kết hợp/sửa lại. Mỗi biến thể có phiên trình duyệt và mốc riêng để kết quả không phụ thuộc biến thể trước. Nếu setup không đủ điều kiện hoặc hậu kiểm không đạt, runner không trả một lượt hoàn chỉnh màu xanh.')
success=None;invalid=None
for row in UI['tests']:
    for checkpoint in row.get('checkpoints',[]):
        if success is None and checkpoint.get('stage')=='save-feedback':success=checkpoint
        if invalid is None and checkpoint.get('stage')=='validation-feedback' and '1.5' in json.dumps(checkpoint):invalid=checkpoint
assert success and invalid
figure(UI_DIR/success['artifacts'][0]['path'],'TC-73 thông báo lưu cấu hình hợp lệ','Bằng chứng CR-E01: thông báo thành công được chụp trước refresh/khôi phục. Vùng trích giữ thông báo và năm trường; checkpoint đi kèm lưu giá trị nhập cùng các giao dịch Credit API.',12.0,(0,0,1416,980))
figure(UI_DIR/invalid['artifacts'][0]['path'],'TC-74 chặn giá thập phân','Bằng chứng CR-E02, vùng trích từ PNG gốc: giá trị 1.5 và thông báo validation xuất hiện cùng nhau. Sau đó tải lại và đọc backend để kiểm tra không lưu cấu hình không hợp lệ.',16.5,(110,330,590,705))
p('Lượt Selenium lưu 140 checkpoint PNG cho 30 biến thể; có 71 hash ảnh khác nhau. Hình giống nhau ở các lần trở về mốc là có thể hợp lý, nên số hash khác nhau không được dùng làm tiêu chí đạt. Điều kiện quan trọng là ảnh được gắn đúng bước, dữ liệu vào/ra đúng và bước bắt buộc đã có bằng chứng.')
p('Kết quả được xác nhận là TC-73 đạt và TC-74 đạt. Mốc thực tế cuối lượt đã được đọc lại là 80/10/5/5/100. Việc kiểm tra thiếu nội dung hoặc chỉ chụp sau teardown không được coi là đủ; bộ evidence đã yêu cầu checkpoint ở thời điểm nhập, có phản hồi và persistence.')
h('3.4. Thực hiện Postman AI outline')
p('Ba request chức năng chạy theo thứ tự thiếu xác thực → prompt rỗng → prompt hợp lệ để xác nhận hai đường từ chối trước khi tiêu hao Credit. Sau đó GET hồ sơ xác minh số dư và GET bài giảng xác minh outline đã lưu. Toàn bộ 15 request và 48 assertion đều đạt trong lượt được chọn; API phản hồi thật, không thay bằng dữ liệu mô phỏng khi nhà cung cấp lỗi.')
figure(OUT/'api-overview.png','Báo cáo API với ba kịch bản và 15 request','Bằng chứng AI-E01: báo cáo phân biệt 3 feature scenarios, 15 request và 48 assertion; mỗi giao dịch có method, URL, HTTP status, trạng thái assertion và độ trễ thực đo.',12.5)
positive=next(r for r in API['requests'] if r['name'].startswith('API-909'))
data=positive['response']['body']['data']
excerpt={'httpStatus':positive['httpStatus'],'success':True,'data':{'outline':{'title':data['outline']['title'],'sections':[data['outline']['sections'][0]]},'creditSpent':data['creditSpent'],'creditBalance':data['creditBalance']}}
p('Trích một phần phản hồi thật của request thành công (chỉ giữ section đầu để dễ đọc; không phải toàn bộ phản hồi):')
code(json.dumps(excerpt,ensure_ascii=False,indent=2))
p('Kết quả dàn ý được đánh giá bằng kiểu dữ liệu và cấu trúc. Số Credit trả về là 5; số dư sau sinh là 10. Hai request GET tiếp theo khớp số dư và outline đã lưu, nên bằng chứng không dừng ở việc “HTTP 200”. Mốc Credit toàn hệ thống được đọc trước/sau và giữ nguyên; suite không có PATCH cấu hình.')
h('3.5. Thu hồi và kiểm tra tính toàn vẹn')
p('TC-9 và API outline dùng namespace user riêng nên có thể thu hồi đúng dữ liệu đã tạo, tránh xóa dữ liệu người dùng có sẵn. Selenium khôi phục cấu hình ban đầu sau lượt chạy. Hai run tự động đều vượt kiểm tra integrity offline trước khi đưa vào báo cáo. Metadata và source snapshot nhận diện mã thực thi, kể cả khi lượt chạy dùng thay đổi chưa commit; không lấy commit cuối cùng để thay thế dấu vết đó.')
p('Các screenshot, JSON phản hồi và credential runtime không được commit công khai. Báo cáo chỉ nhúng bằng chứng chọn lọc đã kiểm tra; không đưa mật khẩu, token, OTP hoặc khóa Gemini vào tài liệu. Bằng chứng gốc được giữ cục bộ trong thư mục lượt chạy và có thể cung cấp riêng khi cần đối chiếu.')

chapter('4. BÁO CÁO KẾT QUẢ KIỂM THỬ')
h('4.1. Kết quả theo ca và chức năng')
caption_table('Kết quả của sáu đơn vị kiểm thử logic',['Chức năng / ca','Kết quả','Căn cứ'],[
 ('Avatar — TC-9','FAILED','PDF bị chặn nhưng khác thông báo; preview PNG vuông thay vì tròn. Lưu, đồng bộ và persistence hoạt động.'),('Credit — TC-73','PASSED','Chuỗi đổi giá hợp lệ, biên 1, giữ giá sau F5 và khôi phục đúng.'),('Credit — TC-74','PASSED','28 biến thể đơn và 1 biến thể kết hợp/sửa lại đều chặn/lưu đúng.'),('Outline — AI-OUTLINE-01','PASSED','200, cấu trúc hợp lệ, đúng charge và đọc lại balance/outline khớp.'),('Outline — AI-OUTLINE-02','PASSED','422 với prompt rỗng; số dư không giảm trước sinh hợp lệ.'),('Outline — AI-OUTLINE-03','PASSED','401 khi không xác thực; có thông báo lỗi.')],[5.0,2.5,8.8])
p('Theo sáu đơn vị logic đã thiết kế, có năm đạt và một không đạt. Đây là mẫu số của sáu mục nêu trên, không phải tỷ lệ bao phủ toàn hệ thống. Ở mức thực thi tự động, Selenium có 30/30 biến thể đạt và Postman có 15/15 request, 48/48 assertion đạt. Không cộng 30 biến thể với 15 request rồi gọi thành 45 test case.')
h('4.2. Các sai lệch phát hiện ở TC-9')
caption_table('Sai lệch so với yêu cầu đã chọn',['Mã nội bộ','Mong đợi và thực tế','Mức độ / tác động'],[
 ('DEF-AV-01','Mong đợi một thông báo chung có “tối đa 2MB”. PDF thực tế dùng thông báo định dạng riêng với danh sách JPG, PNG, WEBP, GIF.','Low đề xuất. Tệp vẫn bị chặn nhưng không khớp nội dung TC; cần thống nhất yêu cầu hoặc sửa thông báo.'),('DEF-AV-02','TC yêu cầu preview trong khung tròn. Giao diện thực tế dùng khung vuông; ảnh và CSS border-radius 0px cùng xác nhận.','Low đề xuất. Ảnh vẫn xem trước/lưu được, nhưng hình thức không đáp ứng yêu cầu.')],[3.0,9.0,4.3])
p('Hai mã trên là mã theo dõi trong báo cáo, không phải issue Jira đã được tạo. Mức độ được đề xuất theo ảnh hưởng quan sát được. Nếu nhóm thống nhất thay đổi yêu cầu cho phù hợp thiết kế mới, phải cập nhật test basis có kiểm soát; không đổi tiêu chí sau chạy chỉ để chuyển kết quả thành PASSED.')
h('4.3. Giới hạn và công việc tiếp theo')
p('Bằng chứng cho phép kết luận về các đường dữ liệu đã chọn, không chứng minh mọi tệp ảnh hoặc mọi prompt đều được xử lý đúng. Mốc đúng bằng 2 MB mới được phân tích, chưa có lượt chạy riêng; TC-9 cũng chưa kiểm tra tệp hỏng/đổi đuôi. Tương tự, các giá Credit cực đại, trình duyệt khác, nhiều Admin đồng thời và UC032-03 chưa được đánh giá trong phạm vi báo cáo.')
p('Một lần outline thành công xác nhận hợp đồng của lần gọi đó, không chứng minh độ ổn định nhiều lần hoặc tính đúng đắn sư phạm. Độ trễ ghi trong report là quan sát của request, không phải kết quả load/performance test. Việc dùng checksum không loại trừ mọi khả năng chỉnh sửa đồng thời manifest hoặc chứng minh người thực hiện độc lập.')
p('Ưu tiên tiếp theo là xử lý hoặc thống nhất hai sai lệch TC-9, sau đó kiểm tra lại đúng các bước bị ảnh hưởng và hồi quy lưu/đồng bộ avatar. Có thể bổ sung bộ dữ liệu quanh biên 2 MB như một phạm vi mới. Không cần mở rộng collection toàn backend trước khi hoàn tất phân tích và khép kín bằng chứng của ba chức năng hiện tại.')
h('4.4. Kết luận')
p('Trong phạm vi đã chọn, hai ca cấu hình Credit và ba kịch bản API outline có bằng chứng đạt. Ca ảnh đại diện thực hiện được các bước chặn tệp sai, lưu và đồng bộ ảnh hợp lệ, nhưng chưa đạt toàn bộ yêu cầu vì khác thông báo PDF và hình dạng preview. Kết luận này giữ cả phần đúng và phần sai của hệ thống để làm căn cứ sửa lỗi, kiểm tra lại và đánh giá chất lượng trong phạm vi bài tập.')

doc.add_page_break();h('PHỤ LỤC',1)
h('A. Danh sách 30 biến thể Selenium')
caption_table('Ma trận thực thi rút gọn của Selenium',['STT','Ca','Biến thể / dữ liệu','Kết quả'],[(str(i+1),'TC-73' if 'uc032_01' in row['test'] else 'TC-74', 'Chuỗi lưu hợp lệ và khôi phục' if 'uc032_01' in row['test'] else (row['test'].split('[')[-1].rstrip(']') if '[' in row['test'] else 'Kết hợp giá sai và sửa lại 6/6'),row['status']) for i,row in enumerate(UI['tests'])],[1.3,2.0,10.0,3.0])
h('B. Danh sách request của collection outline')
inventory=json.loads((QA/'tests/postman/outline/coverage.json').read_text(encoding='utf-8'))
caption_table('Request, vai trò và trạng thái HTTP',['ID','Vai trò','Method và đường dẫn','HTTP'],[(x['id'],x['scenario'] or ('Thiết lập' if i<5 else 'Thu hồi'),x['method']+' '+x['route'],str(x['status'])) for i,x in enumerate(inventory)],[2.0,3.5,9.3,1.5])
doc.add_page_break();h('C. Nhận diện dữ liệu và nguồn bằng chứng')
p('Các hash dưới đây phục vụ đối chiếu tệp, không thay cho việc đọc actual result. Mã fixture đầy đủ trong ảnh là tài khoản thử dùng namespace ngẫu nhiên, không phải thông tin tài khoản cá nhân.')
for x in sorted(FIX,key=lambda x:x['name']):
    p(x['name']+' — '+str(x['bytes'])+' byte',True);code('SHA-256 '+x['sha256'])
p('Giá trị avatar đọc từ GET /auth/me',True);code('SHA-256 '+AV['independentReadback']['avatarValueSha256'])
caption_table('Nguồn hồ sơ thực thi',['Nguồn','Định danh'],[
 ('TC-9',AV['runId']+'; 7 ảnh checkpoint, receipt JSON, dữ liệu tệp và recovery journal.'),('Selenium',UI['runId']+'; summary.json, JUnit, 140 PNG, checkpoint JSON, source snapshot và integrity.json.'),('Postman outline',API['runId']+'; summary.json, JUnit, request/response đã lược bí mật, source snapshot và integrity.json.'),('Mã backend trong hai run tự động',UI['release']['sha']),('Nguồn QA trên GitHub','PR #57: https://github.com/zunohoang/open4um/pull/57 ; phiên bản suite nhỏ được đưa lên nhánh ở commit b863874.'),('Template hình thức','03. TP - Kế hoạch kiểm thử - Group - 03.docx trong Drive BTL_TESTER; chỉ dùng làm chuẩn trình bày.')],[4.0,12.3])
p('Nguồn nội dung: TC-9 do người dùng cung cấp; định nghĩa TC-73/74 trong SPRINT_4_ADMIN_TEST_CASES; hợp đồng backend outline và collection nhỏ; ba hồ sơ thực thi nêu trên. Các hình và trích phản hồi trong chương 3 là một phần của bằng chứng đó. Không đưa trạng thái lịch sử ở một plan Kiwi khác vào kết quả lượt chạy hiện tại.')

# Cached TOC remains a real updateable Word field. Populate entries using verified
# rendered page numbers when available; no dependency on a blank field refresh.
page_map=json.loads((OUT/'toc-pages.json').read_text(encoding='utf-8')) if (OUT/'toc-pages.json').exists() else {}
toc_entries=[]
for para in doc.paragraphs:
    if para.style.name in ['Heading 1','Heading 2'] and para.text!='MỤC LỤC':toc_entries.append((para.text,para.style.name))
for i,(title,style) in enumerate(toc_entries):
    q=toc_marker.insert_paragraph_before()
    q.paragraph_format.space_after=Pt(3);q.paragraph_format.line_spacing=1.0
    q.paragraph_format.left_indent=Cm(0.6 if style=='Heading 2' else 0)
    q.paragraph_format.tab_stops.add_tab_stop(Cm(16.0),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
    if i==0:
        begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');q.add_run()._r.append(begin)
        instruction=OxmlElement('w:instrText');instruction.set(qn('xml:space'),'preserve');instruction.text=' TOC \\o "1-2" \\h \\z \\u '
        q.add_run()._r.append(instruction)
        separate=OxmlElement('w:fldChar');separate.set(qn('w:fldCharType'),'separate');q.add_run()._r.append(separate)
    rr=q.add_run(title+'\t'+str(page_map.get(title,'')));rr.font.size=Pt(11.5);rr.bold=style=='Heading 1'
    if i==len(toc_entries)-1:
        end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');q.add_run()._r.append(end)
toc_marker._element.getparent().remove(toc_marker._element)

# Preserve source section geometry/footer field, strip stale title borders and metadata caches.
for para in doc.paragraphs:
    pr=para._p.pPr
    if pr is not None:
        for child in list(pr):
            if child.tag==qn('w:pBdr'):pr.remove(child)
for name in ['Title','Subtitle']:
    pr=doc.styles[name].element.pPr
    if pr is not None:
        for child in list(pr):
            if child.tag==qn('w:pBdr'):pr.remove(child)
# Source Google Docs editing cache is not formatting and must not carry old data.
for rel in list(doc.part.rels.values()):
    if rel.reltype.endswith('/customXml'):doc.part.drop_rel(rel.rId)
doc.save(OUT/'BTL_TESTER.docx')
assert hashlib.sha256(REF.read_bytes()).hexdigest()==reference_hash
manifest={'referenceSha256':reference_hash,'reportSha256':hashlib.sha256((OUT/'BTL_TESTER.docx').read_bytes()).hexdigest(),
 'seleniumRun':UI['runId'],'apiRun':API['runId'],'avatarRun':AV['runId'],'tables':table_no,'figures':fig_no,
 'caseCounts':{'logical':6,'passed':5,'failed':1,'seleniumVariants':30,'apiRequests':15,'apiAssertions':48}}
(OUT/'report-build.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
print(json.dumps(manifest,indent=2))
