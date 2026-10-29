from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
from docx import Document
from pypdf import PdfReader
import hashlib,json,re,argparse

repo=Path(__file__).resolve().parents[2]
parser=argparse.ArgumentParser(description='Validate report contents, source images and rendered page references offline.')
parser.add_argument('--work-dir',type=Path,default=repo/'tmp/btl-tester')
parser.add_argument('--template',type=Path,required=True)
args=parser.parse_args()
root=args.work_dir.resolve()
reference=args.template.resolve()
docpath=root/'BTL_TESTER.docx'
doc=Document(docpath)
text='\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for row in t.rows for c in row.cells])
chapters=[p.text for p in doc.paragraphs if p.style.name=='Heading 1' and re.match(r'^[1-4]\. ',p.text)]
assert len(chapters)==4 and len(doc.tables)==22
results=next(t for t in doc.tables if t.cell(0,0).text=='Chức năng / ca')
states=[r.cells[1].text for r in results.rows[1:]]
assert states.count('PASSED')==5 and states.count('FAILED')==1
assert 'DEF-AV-01' in text and 'DEF-AV-02' in text and '2.097.152' in text
assert 'Nguyễn Thái Duy' in text and '231210730' in text
assert 'HỌC PHẦN PROJECT 1' not in text and 'TS. Nguyễn Trọng Phúc' not in text
assert not any(x in text for x in ['TODO','TBD','hello','�','Ã¡','Ä‘'])

fixture=json.loads((root/'avatar/summary.json').read_text(encoding='utf-8'))
for step in fixture['steps']:
    f=root/'avatar'/step['screenshot']
    assert hashlib.sha256(f.read_bytes()).hexdigest()==step['screenshotSha256']
assert fixture['headerSynchronized'] and fixture['cleanup']=='verified'

secrets=[]
for line in (repo/'tests/selenium/.env').read_text(encoding='utf-8-sig').splitlines():
    if '=' in line and not line.lstrip().startswith('#'):
        key,value=line.split('=',1)
        if any(w in key.upper() for w in ['PASSWORD','EMAIL','TOKEN','SECRET']):
            value=value.strip().strip('"').strip("'")
            if len(value)>5:secrets.append(value)
with ZipFile(docpath) as z:
    for name in z.namelist():
        if name.endswith('.xml'):
            content=z.read(name).decode('utf-8')
            assert not any(s in content for s in secrets),'Secret detected in document XML'
    assert not any(n.lower().startswith('customxml/') for n in z.namelist())
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    final=E.fromstring(z.read('word/document.xml')).find('.//w:sectPr',ns)
    with ZipFile(reference) as source:
        ref=E.fromstring(source.read('word/document.xml')).find('.//w:sectPr',ns)
        for part in ['pgSz','pgMar']:
            assert dict(final.find('w:'+part,ns).attrib)==dict(ref.find('w:'+part,ns).attrib)
    original_pngs={hashlib.sha256(p.read_bytes()).hexdigest() for p in (root/'avatar').glob('*.png')}
    for p in (repo/'tests/selenium/reports/422a0a7242d44f1c852e96c9115e2031/images').rglob('*.png'):original_pngs.add(hashlib.sha256(p.read_bytes()).hexdigest())
    original_pngs.add(hashlib.sha256((root/'api-overview.png').read_bytes()).hexdigest())
    embedded=[n for n in z.namelist() if n.startswith('word/media/') and n.endswith('.png')]
    assert len(embedded)==6
    for name in embedded:assert hashlib.sha256(z.read(name)).hexdigest() in original_pngs

pdf=PdfReader(root/'render-delivery-2/BTL_TESTER.pdf')
pages=len(pdf.pages)
assert pages==23
norm=lambda s:re.sub(r'\s+',' ',s).strip()
for heading,num in json.loads((root/'toc-pages.json').read_text(encoding='utf-8')).items():
    assert norm(heading) in norm(pdf.pages[num].extract_text() or ''),(heading,num)
assert not any('Error! Reference source not found' in (p.extract_text() or '') for p in pdf.pages)
refhash=hashlib.sha256(reference.read_bytes()).hexdigest()
assert refhash=='eee6ed0cdae50bf3eb21f37209cdbf7a4f5254e2f0cb9fb576daebd803698146'
report={'file':docpath.name,'bytes':docpath.stat().st_size,'sha256':hashlib.sha256(docpath.read_bytes()).hexdigest(),
        'pages':pages,'chapters':4,'tables':22,'evidenceImages':6,'logicalCases':6,'passed':5,'failed':1,
        'checks':['counts and outcomes','source PNG bytes unchanged','TC-9 image hashes','no credentials in XML',
                  'contents references','template geometry','original template unchanged','no stale source Google Docs cache']}
(root/'delivery-validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
